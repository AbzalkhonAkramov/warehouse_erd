from decimal import Decimal, InvalidOperation
from io import BytesIO

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    Response,
    UploadFile,
    status,
)
from openpyxl import Workbook, load_workbook
from sqlalchemy import delete, insert, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.api.deps import get_current_user, require_roles
from app.core.database import get_db
from app.models.associations import customer_agents
from app.models.enums import UserRole
from app.models.sales import Customer, Region
from app.models.user import User
from app.schemas.customer import (
    AgentBrief,
    AgentShopsUpdate,
    CustomerCreate,
    CustomerOut,
    CustomerUpdate,
    VisitDaysUpdate,
)
from app.schemas.imports import BulkImportResult

router = APIRouter(prefix="/customers", tags=["customers"])

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_CUSTOMER_COLUMNS = [
    "Name", "Phone", "Address", "City", "Region", "Credit limit",
    "Visit days (mon,wed,fri)", "Agent (name or email)",
]

_EAGER = (selectinload(Customer.agents), selectinload(Customer.region))


def _to_out(c: Customer) -> CustomerOut:
    out = CustomerOut.model_validate(c)
    out.region_name = c.region.name if c.region else None
    out.agents = [AgentBrief(id=a.id, full_name=a.full_name) for a in c.agents]
    out.agent_ids = [a.id for a in c.agents]
    return out


async def _load(db: AsyncSession, customer_id: int) -> Customer:
    c = await db.scalar(
        select(Customer)
        .where(Customer.id == customer_id)
        .options(*_EAGER)
        .execution_options(populate_existing=True)
    )
    if c is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Customer not found")
    return c


@router.get("", response_model=list[CustomerOut])
async def list_customers(
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> list[CustomerOut]:
    stmt = select(Customer).options(*_EAGER).order_by(Customer.name)
    # Agents only see their own markets; managers/admin see all.
    if user.role == UserRole.AGENT:
        stmt = stmt.where(Customer.agents.any(User.id == user.id))
    return [_to_out(c) for c in await db.scalars(stmt)]


@router.get("/template")
async def customers_template(
    _: User = Depends(require_roles(UserRole.MANAGER)),
) -> Response:
    """Blank market-creation sheet. Region is matched by name (created if new);
    Agent is matched by full name or email. Declared before /{customer_id}."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Markets"
    ws.append(_CUSTOMER_COLUMNS)
    ws.append(["Corner Shop", "+998901234567", "12 Market St", "Tashkent",
               "Center", 500, "mon,wed,fri", "agent@erp.local"])
    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type=_XLSX_MIME,
        headers={"Content-Disposition": 'attachment; filename="markets-template.xlsx"'},
    )


@router.post("/import", response_model=BulkImportResult)
async def customers_import(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_roles(UserRole.MANAGER)),
) -> BulkImportResult:
    """Bulk-create markets. A row needs a Name (unique — existing names skip).
    Region is matched by name (created if new); Agent by full name or email."""
    try:
        wb = load_workbook(BytesIO(await file.read()), data_only=True)
    except Exception:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST, "Could not read the Excel file (.xlsx expected)"
        ) from None
    ws = wb.worksheets[0]

    existing = {(n or "").strip().lower() for n in await db.scalars(select(Customer.name))}
    regions = {(n or "").lower(): rid for rid, n in (
        await db.execute(select(Region.id, Region.name))).all()}
    users = list(await db.scalars(
        select(User).where(User.role == UserRole.AGENT)))
    agent_by = {}
    for u in users:
        agent_by[u.full_name.lower()] = u.id
        agent_by[u.email.lower()] = u.id

    created = skipped = 0
    errors: list[str] = []
    seen: set[str] = set()
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        cell = lambda i: row[i] if row and len(row) > i else None  # noqa: E731
        name = (str(cell(0)).strip() if cell(0) not in (None, "") else "")
        if not name:
            continue
        key = name.lower()
        if key in existing or key in seen:
            skipped += 1
            continue

        region_name = (str(cell(4)).strip() if cell(4) not in (None, "") else "")
        region_id = None
        if region_name:
            region_id = regions.get(region_name.lower())
            if region_id is None:
                r = Region(name=region_name)
                db.add(r)
                await db.flush()
                region_id = r.id
                regions[region_name.lower()] = region_id

        try:
            credit = Decimal(str(cell(5))) if cell(5) not in (None, "") else Decimal("0")
        except (InvalidOperation, ValueError):
            credit = Decimal("0")

        customer = Customer(
            name=name,
            phone=(str(cell(1)).strip() if cell(1) not in (None, "") else None),
            address=(str(cell(2)).strip() if cell(2) not in (None, "") else None),
            city=(str(cell(3)).strip() if cell(3) not in (None, "") else None),
            region_id=region_id,
            credit_limit=credit,
            visit_days=(str(cell(6)).strip() if cell(6) not in (None, "") else None),
        )
        agent_ref = (str(cell(7)).strip() if cell(7) not in (None, "") else "")
        if agent_ref:
            aid = agent_by.get(agent_ref.lower())
            if aid is not None:
                customer.agents = [u for u in users if u.id == aid]
            else:
                errors.append(f"Row {idx}: agent '{agent_ref}' not found (market created without it)")
        db.add(customer)
        seen.add(key)
        created += 1

    await db.flush()
    return BulkImportResult(created=created, skipped=skipped, errors=errors[:50])


@router.get("/{customer_id}", response_model=CustomerOut)
async def get_customer(
    customer_id: int,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> CustomerOut:
    c = await _load(db, customer_id)
    if user.role == UserRole.AGENT and user.id not in [a.id for a in c.agents]:
        raise HTTPException(status.HTTP_403_FORBIDDEN, "This shop is not assigned to you")
    return _to_out(c)


@router.post("", response_model=CustomerOut, status_code=status.HTTP_201_CREATED)
async def create_customer(
    data: CustomerCreate,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> CustomerOut:
    payload = data.model_dump(exclude={"agent_ids"})
    customer = Customer(**payload)

    agent_ids = set(data.agent_ids)
    if user.role == UserRole.AGENT:
        agent_ids.add(user.id)  # an agent creating a shop pins it to themselves
    if agent_ids:
        customer.agents = list(await db.scalars(select(User).where(User.id.in_(agent_ids))))

    db.add(customer)
    await db.flush()
    return _to_out(await _load(db, customer.id))


@router.put("/agent-shops/{agent_id}", response_model=list[CustomerOut])
async def set_agent_shops(
    agent_id: int,
    data: AgentShopsUpdate,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_roles(UserRole.MANAGER)),
) -> list[CustomerOut]:
    """Pin exactly this set of markets to the agent (managers only). Other agents on
    those markets are left untouched (many-to-many)."""
    await db.execute(delete(customer_agents).where(customer_agents.c.agent_id == agent_id))
    for cid in set(data.customer_ids):
        await db.execute(insert(customer_agents).values(customer_id=cid, agent_id=agent_id))
    await db.flush()
    stmt = (
        select(Customer)
        .where(Customer.agents.any(User.id == agent_id))
        .options(*_EAGER)
        .order_by(Customer.name)
    )
    return [_to_out(c) for c in await db.scalars(stmt)]


@router.patch("/{customer_id}", response_model=CustomerOut)
async def update_customer(
    customer_id: int,
    data: CustomerUpdate,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_roles(UserRole.MANAGER)),
) -> CustomerOut:
    customer = await _load(db, customer_id)
    payload = data.model_dump(exclude_unset=True, exclude={"agent_ids"})
    for field, value in payload.items():
        setattr(customer, field, value)
    if data.agent_ids is not None:
        customer.agents = list(
            await db.scalars(select(User).where(User.id.in_(set(data.agent_ids))))
        )
    await db.flush()
    return _to_out(await _load(db, customer.id))


@router.patch("/{customer_id}/visit-days", response_model=CustomerOut)
async def update_visit_days(
    customer_id: int,
    data: VisitDaysUpdate,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_roles(UserRole.MANAGER)),
) -> CustomerOut:
    """Set the market's visit days. Managers/admins only (agents see it read‑only)."""
    customer = await _load(db, customer_id)
    customer.visit_days = data.visit_days
    await db.flush()
    return _to_out(await _load(db, customer.id))
