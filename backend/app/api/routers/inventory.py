from decimal import Decimal, InvalidOperation
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, Response, UploadFile, status
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.deps import get_current_user, require_roles
from app.core.database import get_db
from app.models.catalog import Product, Warehouse
from app.models.enums import StockMovementType, UserRole
from app.models.inventory import Stock, StockMovement
from app.models.user import User

router = APIRouter(prefix="/inventory", tags=["inventory"])

_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class StockOut(BaseModel):
    product_id: int
    product_name: str
    sku: str
    warehouse_id: int
    quantity: Decimal
    min_stock: Decimal
    low: bool


class StockAdjustIn(BaseModel):
    product_id: int
    warehouse_id: int | None = None
    quantity: Decimal = Field(description="Signed delta: +receipt, -adjustment out")
    type: StockMovementType = StockMovementType.ADJUSTMENT
    unit_cost: Decimal = Decimal("0")
    note: str | None = None


async def _default_warehouse_id(db: AsyncSession) -> int:
    wh = await db.scalar(select(Warehouse).where(Warehouse.is_default).limit(1))
    if wh is None:
        wh = await db.scalar(select(Warehouse).order_by(Warehouse.id).limit(1))
    if wh is None:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No warehouse configured")
    return wh.id


@router.get("/stock", response_model=list[StockOut])
async def list_stock(
    low_only: bool = False,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(get_current_user),
) -> list[StockOut]:
    rows = await db.execute(
        select(Stock, Product).join(Product, Product.id == Stock.product_id)
    )
    out: list[StockOut] = []
    for stock, product in rows.all():
        low = Decimal(stock.quantity) <= Decimal(product.min_stock)
        if low_only and not low:
            continue
        out.append(
            StockOut(
                product_id=product.id,
                product_name=product.name,
                sku=product.sku,
                warehouse_id=stock.warehouse_id,
                quantity=Decimal(stock.quantity),
                min_stock=Decimal(product.min_stock),
                low=low,
            )
        )
    return out


@router.post("/adjust", status_code=status.HTTP_201_CREATED)
async def adjust_stock(
    data: StockAdjustIn,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_roles(UserRole.WAREHOUSE, UserRole.MANAGER)),
) -> dict:
    """Manually receive stock or correct a count. Records a StockMovement."""
    warehouse_id = data.warehouse_id or await _default_warehouse_id(db)
    stock = await db.scalar(
        select(Stock).where(
            Stock.product_id == data.product_id, Stock.warehouse_id == warehouse_id
        )
    )
    if stock is None:
        stock = Stock(product_id=data.product_id, warehouse_id=warehouse_id, quantity=Decimal("0"))
        db.add(stock)
    new_qty = Decimal(stock.quantity) + Decimal(data.quantity)
    if new_qty < 0:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "Stock cannot go negative")
    stock.quantity = new_qty
    db.add(
        StockMovement(
            product_id=data.product_id,
            warehouse_id=warehouse_id,
            type=data.type,
            quantity=data.quantity,
            unit_cost=data.unit_cost,
            note=data.note,
            created_by_id=user.id,
        )
    )
    await db.flush()
    return {"product_id": data.product_id, "quantity": float(stock.quantity)}


class StockImportResult(BaseModel):
    updated: int
    added_total: Decimal
    skipped: int
    errors: list[str]


async def _apply_stock_in(
    db: AsyncSession, product_id: int, warehouse_id: int, qty: Decimal, user_id: int
) -> None:
    stock = await db.scalar(
        select(Stock).where(
            Stock.product_id == product_id, Stock.warehouse_id == warehouse_id
        )
    )
    if stock is None:
        stock = Stock(product_id=product_id, warehouse_id=warehouse_id, quantity=Decimal("0"))
        db.add(stock)
    stock.quantity = Decimal(stock.quantity) + qty
    db.add(
        StockMovement(
            product_id=product_id,
            warehouse_id=warehouse_id,
            type=StockMovementType.RECEIPT,
            quantity=qty,
            note="Excel import",
            created_by_id=user_id,
        )
    )


@router.get("/stock-template")
async def stock_template(
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_roles(UserRole.WAREHOUSE, UserRole.MANAGER)),
) -> Response:
    """Download an .xlsx with two sheets: 'Add stock' (id, name, available, blank
    'Add quantity' column D to fill) and 'Guide' (full goods reference list)."""
    warehouse_id = await _default_warehouse_id(db)
    products = list(
        await db.scalars(
            select(Product).where(Product.is_active.is_(True)).order_by(Product.name)
        )
    )
    qty_map = {
        pid: q
        for pid, q in (
            await db.execute(
                select(Stock.product_id, Stock.quantity).where(
                    Stock.warehouse_id == warehouse_id
                )
            )
        ).all()
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "Add stock"
    ws.append(["Product ID", "Name", "Available stock", "Add quantity"])
    for p in products:
        ws.append([p.id, p.name, float(qty_map.get(p.id, 0)), None])

    guide = wb.create_sheet("Guide")
    guide.append(["ID", "Name", "Available stock"])
    for p in products:
        guide.append([p.id, p.name, float(qty_map.get(p.id, 0))])

    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type=_XLSX_MIME,
        headers={"Content-Disposition": 'attachment; filename="stock-template.xlsx"'},
    )


@router.post("/stock-import", response_model=StockImportResult)
async def stock_import(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_roles(UserRole.WAREHOUSE, UserRole.MANAGER)),
) -> StockImportResult:
    """Bulk stock-in: reads the first sheet, adds column D ('Add quantity') to the
    stock of the product in column A. Rows with a blank/zero D are skipped."""
    try:
        wb = load_workbook(BytesIO(await file.read()), data_only=True)
    except Exception:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST, "Could not read the Excel file (.xlsx expected)"
        ) from None
    ws = wb.worksheets[0]
    warehouse_id = await _default_warehouse_id(db)
    valid_ids = set(await db.scalars(select(Product.id)))

    updated = 0
    skipped = 0
    added_total = Decimal("0")
    errors: list[str] = []

    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row:
            continue
        pid_raw = row[0] if len(row) > 0 else None
        add_raw = row[3] if len(row) > 3 else None
        if add_raw in (None, "", 0):
            if pid_raw not in (None, ""):
                skipped += 1
            continue
        try:
            pid = int(pid_raw)
        except (TypeError, ValueError):
            errors.append(f"Row {idx}: invalid product id '{pid_raw}'")
            continue
        if pid not in valid_ids:
            errors.append(f"Row {idx}: product #{pid} not found")
            continue
        try:
            add = Decimal(str(add_raw))
        except (InvalidOperation, ValueError):
            errors.append(f"Row {idx}: invalid quantity '{add_raw}'")
            continue
        if add <= 0:
            skipped += 1
            continue
        await _apply_stock_in(db, pid, warehouse_id, add, user.id)
        updated += 1
        added_total += add

    await db.flush()
    return StockImportResult(
        updated=updated, added_total=added_total, skipped=skipped, errors=errors[:50]
    )
