from fastapi import APIRouter, Depends, HTTPException, Response, status
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

import uuid
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path

from fastapi import File, Form, UploadFile
from openpyxl import Workbook, load_workbook
from sqlalchemy import func

from app.api.deps import get_current_user, require_roles
from app.core.config import settings
from app.core.database import get_db
from app.models.associations import agent_categories
from app.models.enums import SaleMode, SalesOrderStatus, StockMovementType, UserRole
from app.models.catalog import Category, Currency, Product, Warehouse
from app.models.inventory import Stock, StockMovement
from app.models.sales import SalesOrder, SalesOrderLine
from app.models.user import User
from app.schemas.imports import BulkImportResult
from app.schemas.product import (
    ProductCreate,
    ProductHistoryEntry,
    ProductOut,
    ProductUpdate,
)

router = APIRouter(prefix="/products", tags=["products"])

_IMAGE_SUBDIR = "products"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _xlsx_response(wb: Workbook, filename: str) -> Response:
    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type=_XLSX_MIME,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

# Column order shared by the product template and importer.
_PRODUCT_COLUMNS = [
    "SKU", "Name", "Category", "Unit", "Cost price", "Sale price", "Min stock",
    "Currency", "Sale mode (piece/box/both)", "Box qty", "Box weight",
    "Box dimensions", "Initial stock", "Whole units? (yes/no)",
]


def _parse_bool(v: object, default: bool = True) -> bool:
    if v in (None, ""):
        return default
    return str(v).strip().lower() in ("yes", "y", "true", "1", "да", "ha")


@router.get("", response_model=list[ProductOut])
async def list_products(
    search: str | None = None,
    active_only: bool = True,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> list[ProductOut]:
    stmt = select(Product).order_by(Product.name)
    if active_only:
        stmt = stmt.where(Product.is_active.is_(True))
    if search:
        like = f"%{search}%"
        stmt = stmt.where(Product.name.ilike(like) | Product.sku.ilike(like))

    is_agent = user.role == UserRole.AGENT
    # Agents only see products in the categories assigned to them by a manager.
    # No assignment = unrestricted (see everything).
    if is_agent:
        cat_ids = list(
            await db.scalars(
                select(agent_categories.c.category_id).where(
                    agent_categories.c.agent_id == user.id
                )
            )
        )
        if cat_ids:
            stmt = stmt.where(Product.category_id.in_(cat_ids))

    products = list(await db.scalars(stmt))

    # On-hand quantity per product (so agents can see available stock).
    qty_rows = await db.execute(
        select(Stock.product_id, func.coalesce(func.sum(Stock.quantity), 0)).group_by(
            Stock.product_id
        )
    )
    qty_map = {pid: Decimal(q) for pid, q in qty_rows.all()}

    out: list[ProductOut] = []
    for p in products:
        item = ProductOut.model_validate(p)
        item.on_hand = qty_map.get(p.id, Decimal("0"))
        if p.currency is not None:
            item.currency_code = p.currency.code
            item.currency_symbol = p.currency.symbol
        if is_agent:
            item.cost_price = None  # agents never see the purchase price
        out.append(item)
    return out


@router.get("/template")
async def products_template(
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_roles(UserRole.MANAGER, UserRole.WAREHOUSE)),
) -> Response:
    """Blank product-creation sheet plus a 'Guide' sheet listing the categories,
    currency codes and sale modes that are valid to type in. Declared before the
    /{product_id} route so the literal 'template' path isn't captured as an id."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Products"
    ws.append(_PRODUCT_COLUMNS)
    ws.append([
        "SKU-100", "Example product", "Drinks", "шт", 5, 9, 10, "UZS",
        "piece", 24, 7.5, "40x30x25", 100, "yes",
    ])

    guide = wb.create_sheet("Guide")
    cats = list(await db.scalars(select(Category.name).order_by(Category.name)))
    curs = list(await db.scalars(
        select(Currency.code).where(Currency.is_active.is_(True)).order_by(Currency.code)
    ))
    guide.append(["Categories", "Currencies", "Sale modes"])
    for i in range(max(len(cats), len(curs), 3)):
        guide.append([
            cats[i] if i < len(cats) else None,
            curs[i] if i < len(curs) else None,
            ["piece", "box", "both"][i] if i < 3 else None,
        ])
    return _xlsx_response(wb, "products-template.xlsx")


@router.get("/{product_id}", response_model=ProductOut)
async def get_product(
    product_id: int,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> ProductOut:
    product = await db.get(Product, product_id)
    if product is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Product not found")
    total = await db.scalar(
        select(func.coalesce(func.sum(Stock.quantity), 0)).where(Stock.product_id == product_id)
    )
    item = ProductOut.model_validate(product)
    item.on_hand = Decimal(total or 0)
    if product.currency is not None:
        item.currency_code = product.currency.code
        item.currency_symbol = product.currency.symbol
    if user.role == UserRole.AGENT:
        item.cost_price = None
    return item


@router.post("", response_model=ProductOut, status_code=status.HTTP_201_CREATED)
async def create_product(
    data: ProductCreate,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_roles(UserRole.MANAGER, UserRole.WAREHOUSE)),
) -> Product:
    exists = await db.scalar(select(Product).where(Product.sku == data.sku))
    if exists:
        raise HTTPException(status.HTTP_409_CONFLICT, "SKU already exists")
    product = Product(**data.model_dump())
    db.add(product)
    await db.flush()
    return product


@router.patch("/{product_id}", response_model=ProductOut)
async def update_product(
    product_id: int,
    data: ProductUpdate,
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_roles(UserRole.MANAGER, UserRole.WAREHOUSE)),
) -> Product:
    product = await db.get(Product, product_id)
    if product is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Product not found")
    payload = data.model_dump(exclude_unset=True)
    if payload.get("sku") and payload["sku"] != product.sku:
        clash = await db.scalar(
            select(Product).where(Product.sku == payload["sku"], Product.id != product_id)
        )
        if clash:
            raise HTTPException(status.HTTP_409_CONFLICT, "SKU already exists")
    for field, value in payload.items():
        setattr(product, field, value)
    await db.flush()
    return product


def _dec(v: object, default: Decimal | None = Decimal("0")) -> Decimal | None:
    if v in (None, ""):
        return default
    try:
        return Decimal(str(v))
    except (InvalidOperation, ValueError):
        raise ValueError(f"bad number '{v}'") from None


def _int_or_none(v: object) -> int | None:
    if v in (None, ""):
        return None
    return int(float(v))


@router.post("/import", response_model=BulkImportResult)
async def products_import(
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_roles(UserRole.MANAGER, UserRole.WAREHOUSE)),
) -> BulkImportResult:
    """Bulk-create products from the template. A row needs at least SKU + Name.
    Category is matched by name (created if new); currency by code (must exist);
    an 'Initial stock' > 0 is added to the default warehouse."""
    try:
        wb = load_workbook(BytesIO(await file.read()), data_only=True)
    except Exception:
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST, "Could not read the Excel file (.xlsx expected)"
        ) from None
    ws = wb.worksheets[0]

    existing_skus = {s for s in await db.scalars(select(Product.sku))}
    cats = {(n or "").lower(): cid for cid, n in (
        await db.execute(select(Category.id, Category.name))).all()}
    curs = {(c or "").upper(): cid for cid, c in (
        await db.execute(select(Currency.id, Currency.code))).all()}
    base_currency = await db.scalar(
        select(Currency.id).where(Currency.code == "UZS")
    ) or await db.scalar(select(Currency.id).order_by(Currency.id))
    warehouse_id = await db.scalar(
        select(Warehouse.id).where(Warehouse.is_default.is_(True))
    ) or await db.scalar(select(Warehouse.id).order_by(Warehouse.id))
    valid_modes = {m.value for m in SaleMode}

    created = skipped = 0
    errors: list[str] = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        cell = lambda i: row[i] if row and len(row) > i else None  # noqa: E731
        sku = (str(cell(0)).strip() if cell(0) not in (None, "") else "")
        name = (str(cell(1)).strip() if cell(1) not in (None, "") else "")
        if not sku and not name:
            continue
        if not sku or not name:
            errors.append(f"Row {idx}: SKU and Name are both required")
            continue
        if sku in existing_skus:
            skipped += 1
            continue

        # Category — reuse by name, or create it.
        cat_name = (str(cell(2)).strip() if cell(2) not in (None, "") else "")
        category_id = None
        if cat_name:
            category_id = cats.get(cat_name.lower())
            if category_id is None:
                new_cat = Category(name=cat_name)
                db.add(new_cat)
                await db.flush()
                category_id = new_cat.id
                cats[cat_name.lower()] = category_id

        # Currency — must exist; blank falls back to the base currency.
        cur_code = (str(cell(7)).strip().upper() if cell(7) not in (None, "") else "")
        currency_id = curs.get(cur_code, base_currency) if cur_code else base_currency
        if cur_code and cur_code not in curs:
            errors.append(f"Row {idx}: unknown currency '{cur_code}', used default")

        mode = (str(cell(8)).strip().lower() if cell(8) not in (None, "") else "piece")
        if mode not in valid_modes:
            mode = "piece"

        try:
            product = Product(
                sku=sku,
                name=name,
                category_id=category_id,
                unit=(str(cell(3)).strip() if cell(3) not in (None, "") else "шт"),
                cost_price=_dec(cell(4)),
                sale_price=_dec(cell(5)),
                min_stock=_dec(cell(6)),
                currency_id=currency_id,
                sale_mode=mode,
                box_qty=_int_or_none(cell(9)),
                box_weight=_dec(cell(10), None),
                box_dimensions=(str(cell(11)).strip() if cell(11) not in (None, "") else None),
                integer_qty=_parse_bool(cell(13)),
            )
        except ValueError as e:
            errors.append(f"Row {idx}: {e}")
            continue
        db.add(product)
        await db.flush()
        existing_skus.add(sku)
        created += 1

        # Optional opening stock.
        try:
            init = _dec(cell(12), Decimal("0"))
        except ValueError:
            init = Decimal("0")
        if init and init > 0 and warehouse_id is not None:
            db.add(Stock(product_id=product.id, warehouse_id=warehouse_id, quantity=init))
            db.add(StockMovement(
                product_id=product.id,
                warehouse_id=warehouse_id,
                type=StockMovementType.RECEIPT,
                quantity=init,
                note="excel import",
                created_by_id=user.id,
            ))

    await db.flush()
    return BulkImportResult(created=created, skipped=skipped, errors=errors[:50])


@router.post("/{product_id}/image", response_model=ProductOut)
async def upload_product_image(
    product_id: int,
    file: UploadFile = File(...),
    side: str = Form("front"),  # "front" (main) or "back"
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_roles(UserRole.MANAGER, UserRole.WAREHOUSE)),
) -> Product:
    product = await db.get(Product, product_id)
    if product is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Product not found")

    root = Path(settings.UPLOAD_DIR) / _IMAGE_SUBDIR
    root.mkdir(parents=True, exist_ok=True)
    suffix = Path(file.filename or "").suffix.lower() or ".jpg"
    name = f"{uuid.uuid4().hex}{suffix}"
    (root / name).write_bytes(await file.read())

    rel = f"{_IMAGE_SUBDIR}/{name}"
    if side == "back":
        product.image_back_path = rel
    else:
        product.image_path = rel
    await db.flush()
    return product


# Statuses that do NOT represent a real (stock-deducting) sale.
_EXCLUDED_SALE_STATUSES = (
    SalesOrderStatus.NEW,
    SalesOrderStatus.REFUND,
    SalesOrderStatus.CANCELLED,
    # legacy values:
    SalesOrderStatus.REJECTED,
    SalesOrderStatus.DRAFT,
    SalesOrderStatus.PENDING,
)


@router.get("/{product_id}/history", response_model=list[ProductHistoryEntry])
async def product_history(
    product_id: int,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(get_current_user),
) -> list[ProductHistoryEntry]:
    """Movement history for a product.

    Managers/warehouse/accountants see the full history — stock added (by which
    account) and every sale (by which agent). Agents see only their own sales.
    """
    if await db.get(Product, product_id) is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Product not found")

    is_privileged = user.role != UserRole.AGENT
    entries: list[ProductHistoryEntry] = []

    # Stock-in events (receipts / positive adjustments) — privileged roles only.
    if is_privileged:
        rows = await db.execute(
            select(StockMovement, User.full_name)
            .outerjoin(User, User.id == StockMovement.created_by_id)
            .where(StockMovement.product_id == product_id, StockMovement.quantity > 0)
        )
        for mv, account in rows.all():
            detail = f"{mv.type.value} · {mv.note}" if mv.note else mv.type.value
            entries.append(
                ProductHistoryEntry(
                    kind="added",
                    date=mv.created_at,
                    quantity=mv.quantity,
                    user_name=account,
                    detail=detail,
                )
            )

    # Sales — full for privileged roles, own-only for agents.
    sales_stmt = (
        select(SalesOrderLine, SalesOrder, User.full_name)
        .join(SalesOrder, SalesOrder.id == SalesOrderLine.sales_order_id)
        .outerjoin(User, User.id == SalesOrder.agent_id)
        .where(
            SalesOrderLine.product_id == product_id,
            SalesOrder.status.notin_(_EXCLUDED_SALE_STATUSES),
        )
    )
    if not is_privileged:
        sales_stmt = sales_stmt.where(SalesOrder.agent_id == user.id)

    for line, order, agent_name in (await db.execute(sales_stmt)).all():
        entries.append(
            ProductHistoryEntry(
                kind="sale",
                date=order.created_at,
                quantity=line.quantity,
                user_name=agent_name,
                detail=f"#{order.id} · {order.status.value}",
            )
        )

    entries.sort(key=lambda e: e.date, reverse=True)
    return entries
