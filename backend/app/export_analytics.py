"""Export the analytics model to an Excel workbook for Power BI (or any BI tool).

Creates the analytics views (idempotent) and writes one sheet per view to
``uploads/analytics_export.xlsx`` — which, via the Docker bind-mount, appears on
the host at ``backend/uploads/analytics_export.xlsx``.

Run it inside the backend container:

    docker compose exec backend python -m app.export_analytics

Then in Power BI: Get Data -> Excel -> pick that file -> load the sheets.
Re-run this whenever you want fresh numbers, and hit Refresh in Power BI.
"""

import asyncio
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import text

from app.core.database import AsyncSessionLocal

# Sheet name -> view. v_sales_fact is the grain; the rest are convenience rollups.
VIEWS = {
    "SalesFact": "v_sales_fact",
    "AgentKPI": "v_agent_kpi",
    "SalesByDay": "v_sales_by_day",
    "SalesByMonth": "v_sales_by_month",
    "ProductSales": "v_product_sales",
    "CustomerSales": "v_customer_sales",
    "CategorySales": "v_category_sales",
}

_OUT = Path(__file__).resolve().parent.parent / "uploads" / "analytics_export.xlsx"
_DDL = Path(__file__).with_name("analytics_views.sql")


def _cell(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value
    return value


async def run() -> int:
    ddl = _DDL.read_text()
    wb = Workbook()
    wb.remove(wb.active)  # drop the default empty sheet

    async with AsyncSessionLocal() as db:
        # Make sure the views exist / are up to date.
        for statement in (s.strip() for s in ddl.split(";")):
            if statement:
                await db.execute(text(statement))
        await db.commit()

        total_rows = 0
        for sheet_name, view in VIEWS.items():
            result = await db.execute(text(f"SELECT * FROM {view}"))
            columns = list(result.keys())
            rows = result.fetchall()
            ws = wb.create_sheet(title=sheet_name)
            ws.append(columns)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for row in rows:
                ws.append([_cell(v) for v in row])
            ws.freeze_panes = "A2"
            total_rows += len(rows)

    _OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(_OUT)
    print(f"Wrote {_OUT} — {len(VIEWS)} sheets, {total_rows} data rows total.")
    print("On the host this is: backend/uploads/analytics_export.xlsx")
    return 0


def main() -> None:
    raise SystemExit(asyncio.run(run()))


if __name__ == "__main__":
    main()
