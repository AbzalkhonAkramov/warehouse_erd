"""Excel (.xlsx) order receipt / waybill generator.

Modelled on the customer's "Накладные" template: each order is rendered as TWO
identical copies side by side (columns A–E and G–K, a narrow spacer column F),
portrait. Header box (order no/date, client, address, phone, agent, forwarder),
a bordered product table, an "Итог" line, and an "Отпустил / Принял" footer.

Bulk export stacks orders on one sheet with a page break between them, so the
whole selection prints in one go.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.pagebreak import Break

# Receipt labels per Excel language (matches CompanySettings.excel_lang).
LABELS = {
    "ru": {
        "order": "ЗАКАЗ", "from": "от", "client": "Клиент", "address": "Адрес",
        "phone": "Тел", "agent": "Агент", "forwarder": "Экспедитор", "num": "№",
        "name": "Наименование", "qty": "Кол-во", "price": "Цена", "sum": "Сумма",
        "total": "Итог", "currency": "сум", "released": "Отпустил",
        "received": "Принял",
    },
    "uz": {
        "order": "BUYURTMA", "from": "sana", "client": "Mijoz", "address": "Manzil",
        "phone": "Tel", "agent": "Agent", "forwarder": "Ekspeditor", "num": "№",
        "name": "Nomi", "qty": "Soni", "price": "Narx", "sum": "Summa",
        "total": "Jami", "currency": "so‘m", "released": "Berdi",
        "received": "Oldi",
    },
    "en": {
        "order": "ORDER", "from": "of", "client": "Client", "address": "Address",
        "phone": "Tel", "agent": "Agent", "forwarder": "Forwarder", "num": "#",
        "name": "Name", "qty": "Qty", "price": "Price", "sum": "Sum",
        "total": "Total", "currency": "", "released": "Released by",
        "received": "Received by",
    },
}

# One receipt copy is 5 columns wide; a narrow spacer separates the two copies.
_COPY_COLS = 5
_GAP = 1
_WIDTHS = [3, 20, 7, 8, 11]  # №, name, qty, price, sum
_GAP_WIDTH = 4
_MONEY_FMT = "#,##0"
_QTY_FMT = "#,##0.###"

_FONT = Font(name="Calibri", size=8)
_FONT_B = Font(name="Calibri", size=8, bold=True)
_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

_THIN = Side(style="thin", color="000000")
_MED = Side(style="medium", color="000000")
# Order (product table) cells: bold black box. Meta/footer: thin black.
_GRID = Border(left=_MED, right=_MED, top=_MED, bottom=_MED)
_SIDES = Border(left=_THIN, right=_THIN)
_SIDES_TOP = Border(left=_THIN, right=_THIN, top=_THIN)
_BOTTOM = Border(bottom=_THIN)


def _write_copy(ws, r0: int, c0: int, order: dict, L: dict) -> int:
    """Write one receipt copy with its top-left at (r0, c0). Returns the last row used."""
    cnum, cname, cqty, cprice, csum = c0, c0 + 1, c0 + 2, c0 + 3, c0 + 4

    def put(r, c, value, font=_FONT, align=_left, fmt=None, border=None):
        cc = ws.cell(row=r, column=c, value=value)
        cc.font = font
        cc.alignment = align
        if fmt:
            cc.number_format = fmt
        if border:
            cc.border = border
        return cc

    def meta(r, text, bold=False, top=False):
        ws.merge_cells(start_row=r, start_column=cnum, end_row=r, end_column=csum)
        put(r, cnum, text, font=_FONT_B if bold else _FONT, align=_left)
        # Draw the left/right box sides across the merged row (+ top on the first row).
        for c in range(cnum, csum + 1):
            edge = Border(
                left=_THIN if c == cnum else None,
                right=_THIN if c == csum else None,
                top=_THIN if top else None,
            )
            ws.cell(row=r, column=c).border = edge

    r = r0
    meta(r, f'{L["order"]} №{order["order_no"]} {L["from"]} {order["date"]}', bold=True, top=True)
    r += 1
    meta(r, f'{L["client"]}: {order["customer_name"]}')
    r += 1
    meta(r, f'{L["address"]}: {order.get("address") or "—"}')
    r += 1
    meta(r, f'{L["phone"]}: {order.get("phone") or ""}')
    r += 1
    meta(r, f'{L["agent"]}: {order["agent_name"]}')
    r += 1
    meta(r, f'{L["forwarder"]}: {order.get("deliverer") or ""}')
    r += 1

    # Table header (bold black box on every order cell).
    put(r, cnum, L["num"], font=_FONT_B, align=_center, border=_GRID)
    put(r, cname, L["name"], font=_FONT_B, align=_center, border=_GRID)
    put(r, cqty, L["qty"], font=_FONT_B, align=_center, border=_GRID)
    put(r, cprice, L["price"], font=_FONT_B, align=_center, border=_GRID)
    put(r, csum, L["sum"], font=_FONT_B, align=_center, border=_GRID)
    r += 1

    for i, ln in enumerate(order["lines"], start=1):
        put(r, cnum, i, align=_center, border=_GRID)
        put(r, cname, ln["name"], align=_left, border=_GRID)
        put(r, cqty, ln["quantity"], align=_center, fmt=_QTY_FMT, border=_GRID)
        put(r, cprice, ln["unit_price"], align=_right, fmt=_MONEY_FMT, border=_GRID)
        put(r, csum, ln["line_total"], align=_right, fmt=_MONEY_FMT, border=_GRID)
        ws.row_dimensions[r].height = 28
        r += 1

    # Итог (total) row.
    put(r, cnum, None, border=_GRID)
    put(r, cname, L["total"], font=_FONT_B, align=_left, border=_GRID)
    put(r, cqty, order["total_qty"], font=_FONT_B, align=_center, fmt=_QTY_FMT, border=_GRID)
    put(r, cprice, None, border=_GRID)
    cur = f' {L["currency"]}'.rstrip()
    put(r, csum, f'{order["total"]:,.0f}{cur}', font=_FONT_B, align=_right, border=_GRID)
    r += 1

    # Footer: signatures.
    ws.merge_cells(start_row=r, start_column=cname, end_row=r, end_column=cqty)
    put(r, cname, f'{L["released"]}: _______________', align=_left, border=_BOTTOM)
    ws.cell(row=r, column=cqty).border = _BOTTOM
    ws.merge_cells(start_row=r, start_column=cprice, end_row=r, end_column=csum)
    put(r, cprice, f'{L["received"]}: _______________', align=_left, border=_BOTTOM)
    ws.cell(row=r, column=csum).border = _BOTTOM
    ws.row_dimensions[r].height = 25
    return r


def _order_block(ws, start_row: int, order: dict, L: dict) -> int:
    r1 = _write_copy(ws, start_row, 1, order, L)
    r2 = _write_copy(ws, start_row, 1 + _COPY_COLS + _GAP, order, L)
    return max(r1, r2)


def build_receipts_xlsx(orders: list[dict], company: str, lang: str = "ru") -> bytes:
    L = LABELS.get(lang, LABELS["ru"])

    wb = Workbook()
    ws = wb.active
    ws.title = "Receipts"
    ws.sheet_view.showGridLines = False  # others white: only the order box shows

    def set_widths(base_col: int):
        for i, w in enumerate(_WIDTHS):
            ws.column_dimensions[ws.cell(row=1, column=base_col + i).column_letter].width = w

    set_widths(1)
    ws.column_dimensions[ws.cell(row=1, column=1 + _COPY_COLS).column_letter].width = _GAP_WIDTH
    set_widths(1 + _COPY_COLS + _GAP)

    row = 1
    for idx, order in enumerate(orders):
        last = _order_block(ws, row, order, L)
        if idx < len(orders) - 1:
            ws.row_breaks.append(Break(id=last + 1))  # page break between orders
        row = last + 3

    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = ws.page_margins.right = 0.2
    ws.page_margins.top = 0.2
    ws.page_margins.bottom = 0.1
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
